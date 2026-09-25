#!/usr/bin/env python3
# check_monster_sheet.py — keeps the Zone Spawn Sheet (~/NCT/Aldenexia-Lightfall/zones/Zone Spawn Sheet.xlsx) and
# Data/monsters.json matching. monsters.json is the source of truth for every monster; the sheet is the designer's view
# of each zone's monsters (one tab per zone). The first column of each tab is the monster's id in monsters.json.
#
#   python3 tools/check_monster_sheet.py          report differences (exit 1 if any)
#   python3 tools/check_monster_sheet.py --write  make the sheet match monsters.json: ids in the first column, level ranges,
#                                                 and a row for every monster that spawns in the zone but has none
#                                                 (a backup "Zone Spawn Sheet.before_sync_<date>.xlsx" is kept)
#
# Checked per tab: every row names a monster that exists; its level range equals level_min-level_max; its Area lists the
# places it spawns (the spawn points' spawn_zone labels); every monster that SPAWNS in that zone (Data/<zone>_spawns.json)
# has a row. Tabs for zones not built yet are only checked for their ids.
import datetime
import json
import os
import shutil
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET = os.path.expanduser("~/NCT/Aldenexia-Lightfall/zones/Zone Spawn Sheet.xlsx")
TABS = {"LumoraOutskirts": "lumora_outskirts", "DustwindPlateaus": "dustwind_plateaus", "AshfallDunes": "ashfall_dunes"}
ID_HEADER = "Monster ID (monsters.json)"


def load_json(path):
    with open(os.path.join(ROOT, path)) as f:
        return json.load(f)


def areas_in(zone):
    """monster id -> "Stone Circles, Mirage Zone" (the spawn points' labels, tidied), for the sheet's Area column."""
    path = os.path.join(ROOT, "Data", "%s_spawns.json" % zone)
    try:
        data = load_json(path)
    except (OSError, json.JSONDecodeError):
        return {}
    out = {}
    for sp in data.get("spawns", []):
        label = str(sp.get("spawn_zone", ""))
        for suffix in ("_named", "_rare", "_north", "_south", "_east", "_west"):
            label = label[: -len(suffix)] if label.endswith(suffix) else label
        label = label.replace("_", " ").title()
        names = out.setdefault(sp.get("mob_type"), [])
        if label and label not in names:
            names.append(label)
    return {k: ", ".join(v) for k, v in out.items()}


def spawned_in(zone):
    path = os.path.join(ROOT, "Data", "%s_spawns.json" % zone)
    if not os.path.exists(path):
        return []
    try:
        data = load_json(path)
    except json.JSONDecodeError:
        return []
    seen = []
    for s in data.get("spawns", []):
        if s.get("mob_type") not in seen:
            seen.append(s.get("mob_type"))
    return seen


def level_text(mon):
    lo, hi = mon.get("level_min", mon.get("level")), mon.get("level_max", mon.get("level"))
    return str(lo) if lo == hi else "%s-%s" % (lo, hi)


def norm_level(v):
    return str(v).replace(" ", "").replace("–", "-") if v is not None else ""


def main():
    write = "--write" in sys.argv
    if not os.path.exists(SHEET):
        print("No sheet at %s — nothing to check." % SHEET)
        return 0
    monsters = {k: v for k, v in load_json("Data/monsters.json").items() if k != "_comment"}
    wb = openpyxl.load_workbook(SHEET)
    problems = []
    notes = []
    changed = False
    for tab, zone in TABS.items():
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        built = os.path.exists(os.path.join(ROOT, "Data", "%s_spawns.json" % zone))
        headers = [c.value for c in ws[1]]
        col_level = headers.index("Level") + 1 if "Level" in headers else 3
        rows = {}
        for r in range(2, ws.max_row + 1):
            mid = ws.cell(r, 1).value
            if not any(ws.cell(r, c).value for c in range(1, ws.max_column + 1)):
                continue
            mid = str(mid).strip() if mid else ""
            if mid not in monsters:
                if built:
                    problems.append("%s row %d: '%s' is not a monster in monsters.json" % (tab, r, mid or "(no id)"))
                else:
                    notes.append("%s row %d: '%s' (zone not built yet — add it to monsters.json when it is)" % (tab, r, mid or "(no id)"))
                continue
            rows[mid] = r
            if os.path.exists(os.path.join(ROOT, "Data", "%s_spawns.json" % zone)) and monsters[mid].get("zone") in (zone, None):
                want = level_text(monsters[mid])
                if norm_level(ws.cell(r, col_level).value) != want:
                    problems.append("%s %s: sheet level %s, monsters.json %s" % (tab, mid, ws.cell(r, col_level).value, want))
                    if write:
                        ws.cell(r, col_level).value = want
                        changed = True
        col_area = headers.index("Area") + 1 if "Area" in headers else 0
        areas = areas_in(zone) if built else {}
        for mid, r in rows.items():
            if col_area and mid in areas and str(ws.cell(r, col_area).value or "") != areas[mid]:
                problems.append("%s %s: sheet Area '%s', spawns '%s'" % (tab, mid, ws.cell(r, col_area).value or "", areas[mid]))
                if write:
                    ws.cell(r, col_area).value = areas[mid]
                    changed = True
        for mid in spawned_in(zone):
            if mid in monsters and mid not in rows:
                problems.append("%s: %s spawns there but has no row" % (tab, mid))
                if write:
                    mon = monsters[mid]
                    r = ws.max_row + 1
                    ws.cell(r, 1).value = mid
                    ws.cell(r, 2).value = str(mon.get("description", mid))
                    ws.cell(r, col_level).value = level_text(mon)
                    ws.cell(r, 4).value = str(mon.get("secret_note", "")) or "(added from monsters.json)"
                    ws.cell(r, 5).value = str(mon.get("category", "")).capitalize()
                    ws.cell(r, 6).value = str(mon.get("special_ability", ""))
                    ws.cell(r, 7).value = "" if mon.get("faction") in (None, "None") else mon.get("faction")
                    changed = True
        if write and ws.cell(1, 1).value != ID_HEADER:
            ws.cell(1, 1).value = ID_HEADER
            changed = True
    if notes:
        print("%d row(s) on tabs for zones not built yet aren't in monsters.json (fine for now)." % len(notes))
        if "-v" in sys.argv:
            for n in notes:
                print("   . " + n)
    for p in problems:
        print(" - " + p)
    if not problems:
        print("Zone Spawn Sheet matches monsters.json.")
    if write and changed:
        backup = SHEET.replace(".xlsx", ".before_sync_%s.xlsx" % datetime.date.today().isoformat())
        if not os.path.exists(backup):
            shutil.copy2(SHEET, backup)
        wb.save(SHEET)
        print("Sheet updated (backup: %s)." % os.path.basename(backup))
    return 1 if problems and not write else 0


if __name__ == "__main__":
    sys.exit(main())
