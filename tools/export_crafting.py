#!/usr/bin/env python3
# export_crafting.py — turns the crafting design workbook (Crafting.xlsx) into the game's data files.
#
#   python3 tools/export_crafting.py [path/to/Crafting.xlsx]
#
# Writes:
#   Data/tradeskill_recipes.json  every crafting recipe, keyed by recipe_id
#   Data/gathering_nodes.json     every gatherable node type (Forage / Prospecting / Woodworking harvesting)
#   Data/items.json               APPENDS a definition for every new material, crafted item, kit, tool and recipe
#                                 scroll the workbook needs; entries already in items.json are never changed
#
# The workbook is the source of truth for recipes and nodes: edit it, re-run this, never hand-edit the first two files.
# Item properties (stats, value, description, icon) live in Data/items.json once an item exists: edit them there.
# Anything in an "Effect / Stats" cell the engine can't do yet is listed in the report at the end.

import json
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.expanduser("~/NCT/Aldenexia-Lightfall/Crafting.xlsx")
ICON_DIR = os.path.join(ROOT, "Assets", "icons", "items")

CRAFT_SKILLS = ["Cooking", "Alchemy", "Brewing", "Tinkering", "Blacksmithing", "Woodworking", "Fletching",
                "Jewelcrafting", "Leatherworking", "Tailoring"]
GATHER_TABS = {"Forage": "forage", "Prospecting": "prospecting", "Woodworking": "woodworking"}

# Engine stat names (combatnode.gd gear_<stat>); the sheet's Stamina/Agility map onto them.
STAT_MAP = {"strength": "strength", "dexterity": "dexterity", "agility": "dexterity", "stamina": "constitution",
            "constitution": "constitution", "intelligence": "intelligence", "wisdom": "wisdom", "charisma": "charisma"}
ALL_STATS = ["strength", "constitution", "dexterity", "intelligence", "wisdom"]

SLOT_MAP = {"primary": "primary", "two-handed": "primary", "ranged": "ranged", "offhand shield": "offhand",
            "chest": "chest", "feet": "feet", "hands": "hands", "wrist": "wrist", "finger": "finger", "neck": "neck",
            "charm": "charm", "ammo slot": "ammo", "light slot": "light", "head": "head", "shoulders": "shoulders",
            "arms": "arms", "waist": "waist", "legs": "legs"}

# Placeholder icons until the dedicated art exists (Assets/icons/items/<item_id>.png is used automatically once it does).
PLACEHOLDER_BY_ID = {
    "tin_ore": "ore.png", "copper_ore": "copper-ore.png", "flint": "basalt-chunk.png", "uncut_gem": "gemstone.png",
    "polished_gem": "gemstone.png", "gem_setting": "gemstone.png", "oasis_date": "food.png", "raw_hide": "animal-pelt.png",
    "tanning_oil": "flask.png", "crystal_vial": "flask.png", "dye": "reagent.png", "soft_leather": "leather.png",
    "hardened_leather": "leather.png", "reinforced_leather": "leather.png", "thread": "spider-web.png",
    "woven_cloth": "cloth.png", "reinforced_cloth": "cloth.png", "silken_cloth": "cloth.png",
    "tin_ingot": "ore.png", "copper_ingot": "copper-ore.png", "bronze_ingot": "ore.png",
    "fir_log": "totemfragment.png", "ironwood_log": "totemfragment.png", "palm_log": "totemfragment.png",
    "fir_planks": "totemfragment.png", "ironwood_planks": "totemfragment.png", "palm_planks": "totemfragment.png",
    "fletching_feathers": "bat-wing.png", "tinkered_spring": "scrap_metal.png", "bowstring": "spider-web.png",
    "silk_bowstring": "spider-web.png", "broadhead_arrowheads": "basalt-chunk.png",
    "desert_thistle": "herb.png", "sandroot": "herb.png", "venom_cactus_flower": "herb.png", "moonpetal_sage": "herb.png",
    "tin_dagger": "dagger.png", "copper_war_axe": "axe.png", "copper_armor_plating": "plate-breastplate.png",
    "bronze_shield": "shield.png", "leather_boots": "leather-boots.png", "leather_bracers": "leather-bracers.png",
    "copper_wrist_cuff": "chain-bracers.png", "cloth_gloves": "cloth-bracers.png", "bronze_lockpick_set": "key.png",
    "copper_pickaxe": "axe.png", "bronze_woodcutters_axe": "axe.png", "miners_pick": "axe.png",
    "woodcutters_axe": "axe.png", "smoke_grenade": "pouch.png", "explosive_device": "pouch.png",
    "copper_lantern": "torch.png", "bronze_charm": "cog_trinket.png",
}
PLACEHOLDER_BY_SLOT = {"primary": "sword.png", "ranged": "totemfragment.png", "chest": "leather-chestpiece.png",
                       "finger": "gemstone.png", "neck": "gemstone.png", "ammo": "pouch.png", "light": "torch.png",
                       "head": "hood.png", "shoulders": "tunic.png", "arms": "cloth-bracers.png", "waist": "leather.png",
                       "legs": "cloth_pants.png", "hands": "cloth_gloves.png", "feet": "cloth-shoes.png", "wrist": "cloth-bracers.png"}

# Items that only come from vendors (not made by any recipe or node) — kits, tools, supplies, vendor ammo.
VENDOR_ITEMS = {
    "tanning_oil": ("Tanning Oil", "misc", 3, "A stoppered flask of oil for curing raw hides into leather.", {}),
    "dye": ("Dye", "misc", 3, "A small pot of desert dye, ground from ochre and cactus pulp.", {}),
    "crystal_vial": ("Crystal Vial", "misc", 2, "An empty glass vial. Every potion and poison needs one.", {}),
    "miners_pick": ("Miner's Pick", "tool", 20, "A plain iron pick. Needed to mine ore veins. A tinker can upgrade it.",
                    {"gather_tool": "prospecting", "tool_tier": 1}),
    "woodcutters_axe": ("Woodcutter's Axe", "tool", 20, "A plain iron hatchet. Needed to cut trees. A tinker can upgrade it.",
                        {"gather_tool": "woodworking", "tool_tier": 1}),
    "quiver_of_rough_arrows": ("Quiver of Rough Arrows", "armor", 15,
                               "A quiver of rough arrows you recover and re-fletch as you go. It never runs out.",
                               {"slot": "ammo", "ranged_damage_bonus": 1}),
    "case_of_rough_bolts": ("Case of Rough Bolts", "armor", 15,
                            "A case of rough crossbow bolts you recover as you go. It never runs out.",
                            {"slot": "ammo", "ranged_damage_bonus": 1}),
}
KIT_NAMES = {  # station_id -> skill, for the portable kits sold by vendors
    "basic_cooking_kit": "cooking", "basic_brewing_kit": "brewing", "basic_fletching_kit": "fletching",
    "basic_jewelcrafting_kit": "jewelcrafting", "basic_leatherworking_kit": "leatherworking",
    "basic_tailoring_kit": "tailoring", "basic_woodworking_kit": "woodworking",
    "basic_alchemy_kit": "alchemy", "basic_tinkering_kit": "tinkering",
}
KIT_ICONS = {"basic_cooking_kit": "pouch.png", "basic_brewing_kit": "goblet.png", "basic_fletching_kit": "pouch.png",
             "basic_jewelcrafting_kit": "pouch.png", "basic_leatherworking_kit": "leather.png",
             "basic_tailoring_kit": "cloth.png", "basic_woodworking_kit": "trunk.png"}
# Raw meats: dropped by mobs, not made by a recipe
MEATS = {"raw_spider_meat": "spider", "raw_rat_meat": "rat", "raw_bat_meat": "bat", "raw_goblin_meat": "goblin",
         "raw_human_meat": "human", "raw_snake_meat": "snake"}

report = []


def rows(ws, header_row=1):
    """Yield each table row as a dict keyed by header, stopping at the first blank Recipe ID."""
    headers = [c.value for c in ws[header_row]]
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not r[0]:
            break
        yield dict(zip(headers, r))


def pct(v):
    """'95%' / 0.95 / 95 -> 0.95"""
    if isinstance(v, str):
        v = float(v.strip().rstrip("%")) / 100.0
    elif v > 1:
        v = v / 100.0
    return round(float(v), 4)


def parse_ingredients(text):
    out = {}
    for part in str(text or "").split(","):
        part = part.strip()
        if not part:
            continue
        item_id, qty = part.split(":")
        out[item_id.strip()] = int(qty)
    return out


def parse_cost(text):
    """'20 silver' -> (20, 'silver'); blank / '—' -> (0, 'copper')"""
    m = re.match(r"\s*(\d+)\s*(copper|silver|gold)", str(text or ""), re.I)
    return (int(m.group(1)), m.group(2).lower()) if m else (0, "copper")


def minutes_or_seconds(text):
    m = re.search(r"for (\d+) (min|sec)", text)
    if not m:
        return 0.0
    return float(m.group(1)) * (60.0 if m.group(2) == "min" else 1.0)


def icon_for(item_id, slot="none"):
    if os.path.exists(os.path.join(ICON_DIR, item_id + ".png")):
        return f"res://Assets/icons/items/{item_id}.png"
    name = PLACEHOLDER_BY_ID.get(item_id) or PLACEHOLDER_BY_SLOT.get(slot)
    if not name:
        name = "food.png" if item_id.startswith("raw_") or item_id.startswith("cooked_") else "reagent.png"
    return f"res://Assets/icons/items/{name}"


def base_item(item_id, name, description, item_type="misc", value=2):
    return {
        "name": name, "icon": icon_for(item_id), "description": description, "stackable": True,
        "currency_type": "copper", "value": value, "type": item_type, "armor_class": 0, "slot": "none",
        "skill": "none", "damage": 0, "delay": 0, "weight": 0.3, "size": "small", "class": ["all"],
        "race": ["all"], "ratio": 0, "range": 0, "stat_modifiers": None, "effect": None,
    }


def add_stats(item, text):
    for amount, stat in re.findall(r"\+(\d+) (Strength|Dexterity|Agility|Stamina|Intelligence|Wisdom)\b(?! for)", text):
        mods = item["stat_modifiers"] or {}
        key = STAT_MAP[stat.lower()]
        mods[key] = mods.get(key, 0) + int(amount)
        item["stat_modifiers"] = mods


def apply_effect(item, item_id, effect, unsupported):
    """Turn a sheet 'Effect / Stats' cell into engine item fields. Anything not understood goes to `unsupported`."""
    text = str(effect or "").strip()
    m = re.match(r"^(Primary|Two-handed|Ranged|Offhand shield|Chest|Feet|Hands|Wrist|Finger|Neck|Charm|Ammo slot|Light slot|Head|Shoulders|Arms|Waist|Legs):\s*(.*)$", text, re.I)
    if m:
        slot = SLOT_MAP[m.group(1).lower()]
        rest = m.group(2)
        item["slot"], item["stackable"], item["icon"] = slot, False, icon_for(item_id, slot)
        item["type"] = "armor"
        known = []
        w = re.search(r"(Slashing|Piercing|Blunt)?,?\s*Dmg (\d+) / Delay (\d+)", rest, re.I)
        if w:
            item["type"] = "weapon"
            item["damage"], item["delay"] = int(w.group(2)), int(w.group(3))
            item["skill"] = "archery" if slot == "ranged" else f"{w.group(1).lower()}_weapons"
            item["weight"] = 3.0
            known.append(w.group(0))
        if m.group(1).lower() == "two-handed":
            item["two_handed"] = True
        ac = re.search(r"AC (\d+)", rest)
        if ac:
            item["armor_class"] = int(ac.group(1))
            item["weight"] = 2.0
            known.append(ac.group(0))
        add_stats(item, rest)
        known += re.findall(r"\+\d+ (?:Strength|Dexterity|Agility|Stamina|Intelligence|Wisdom)\b", rest)
        rd = re.search(r"\+(\d+) ranged damage", rest)
        if rd:
            item["ranged_damage_bonus"] = int(rd.group(1))
            known.append(rd.group(0))
        if slot == "light":
            radius = re.search(r"radius (\d+) m", rest)
            burn = re.search(r"burns (\d+) min", rest)
            item["light_source"] = {
                "kind": "fire", "radius": float(radius.group(1)) if radius else 9.0, "energy": 2.2,
                "color": [1.0, 0.72, 0.35], "burn_minutes": float(burn.group(1)) if burn else 0.0,
                "never_burns_out": "never burns out" in rest, "snuffed_by_rain": "stays lit in rain" not in rest,
                "breaks_stealth": "breaks stealth" in rest,
                "light_message": f"You light your {item['name'].lower()}.",
                "rain_snuff_message": f"The rain snuffs out your {item['name'].lower()}.",
            }
            item["stackable"] = "never burns out" not in rest
            return
        leftovers = rest
        for k in known:
            leftovers = leftovers.replace(k, "")
        leftovers = re.sub(r"(Slashing|Piercing|Blunt|Permanent|Needs a quiver in the ammo slot|\(.*?\)|[,.\s])+", " ", leftovers).strip()
        if leftovers:
            unsupported.append(leftovers)
        return

    if re.match(r"^(Material|Component)", text, re.I):
        return
    if text.lower().startswith("tool upgrade") or text.lower().startswith("tool:"):
        item["type"], item["stackable"] = "tool", False
        if "mining" in text:
            item.update({"gather_tool": "prospecting", "tool_tier": 2})
        elif "chopping" in text:
            item.update({"gather_tool": "woodworking", "tool_tier": 2})
        else:
            unsupported.append(text)
        speed = re.search(r"\+(\d+)% \w+ speed", text)
        crit = re.search(r"\+(\d+)% \w+ crit", text)
        if speed:
            item["gather_speed_bonus"] = int(speed.group(1)) / 100.0
        if crit:
            item["gather_crit_bonus"] = int(crit.group(1)) / 100.0
        return
    if text.lower().startswith("use to light a campfire") or text.lower().startswith("throw:"):
        item["type"] = "consumable"
        unsupported.append(text)
        return

    # Consumables: food, drink, potions, poisons
    poison = re.search(r"\+(\d+) weapon damage on hit for (\d+) min", text)
    if poison:
        item["type"] = "consumable"
        item["weapon_poison_bonus_damage"] = int(poison.group(1))
        item["weapon_poison_duration"] = int(poison.group(2)) * 60
        return
    sat = re.search(r"\+(\d+) Satiety", text)
    thirst = re.search(r"\+(\d+) Thirst", text)
    if sat:
        item.update({"type": "food", "restores": "satiety", "restore_amount": int(sat.group(1))})
    elif thirst:
        item.update({"type": "drink", "restores": "thirst", "restore_amount": int(thirst.group(1))})
    else:
        item["type"] = "potion"
    heal = re.search(r"(?:heals|Restores) (\d+) Health", text, re.I)
    if heal:
        item["heal_amount"] = int(heal.group(1))
    if "Removes all active poison" in text:
        item["cures_poison"] = True

    buff_mods, tick_heal = {}, 0
    duration = minutes_or_seconds(text)
    for amount in re.findall(r"\+(\d+) HP/Mana/Stamina regen", text):
        for k in ("hp_regen_bonus", "mana_regen_bonus", "stamina_regen_bonus"):
            buff_mods[k] = int(amount)
    hs = re.search(r"\+(\d+) Health and \+(\d+) Stamina regen", text)
    if hs:
        buff_mods["hp_regen_bonus"], buff_mods["stamina_regen_bonus"] = int(hs.group(1)), int(hs.group(2))
    else:
        st = re.search(r"\+(\d+) Stamina regen", text)
        if st and "HP/Mana/Stamina" not in text:
            buff_mods["stamina_regen_bonus"] = int(st.group(1))
    tick = re.search(r"\+\s?(\d+) Health per tick", text)
    if tick:
        tick_heal = int(tick.group(1))
    for amount, stat in re.findall(r"\+(\d+) (Strength|Dexterity|Agility|Stamina|Intelligence|Wisdom)(?= for| and|,)", text):
        buff_mods["stat_" + STAT_MAP[stat.lower()]] = int(amount)
    allstats = re.search(r"\+(\d+) all stats", text)
    if allstats:
        for s in ALL_STATS:
            buff_mods["stat_" + s] = int(allstats.group(1))
    dmg = re.search(r"\+(\d+)% damage", text)
    if dmg:
        buff_mods["damage_mult"] = int(dmg.group(1)) / 100.0
    if buff_mods or tick_heal:
        item["use_buff"] = {"effect_name": item_id, "duration": duration or 600.0, "modifiers": buff_mods,
                            "tick_heal": tick_heal, "tick_interval": 6.0}
    for phrase in ("weapon damage for", "faster attacks", "elemental resists", "poison immunity", "Perception"):
        if phrase in text:
            unsupported.append(phrase)


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = load_workbook(xlsx, data_only=True)
    items_path = os.path.join(ROOT, "Data", "items.json")
    with open(items_path) as f:
        hand_items = json.load(f)

    items, recipes, nodes = {}, {}, {}

    # Materials tab: every raw / refined material
    for m in rows(wb["Materials"]):
        item_id = m["Item ID"]
        if item_id in hand_items:
            continue
        items[item_id] = base_item(item_id, m["Name"], f"{m['Category']}. {m['Where']}.")

    # Gathering tabs (Woodworking's harvesting table sits under its recipes)
    for tab, skill in GATHER_TABS.items():
        ws = wb[tab]
        header_row = 1
        if tab == "Woodworking":
            header_row = next(c.row for c in ws["A"] if c.value == "Recipe ID" and c.row > 1)
        for g in rows(ws, header_row):
            node_id = re.sub(r"[^a-z0-9]+", "_", g["Node"].lower()).strip("_")
            bonus = None
            b = re.match(r"(.+?) \((\d+)%\)", str(g.get("Bonus Drop") or ""))
            if b:
                bonus_id = next((k for k, v in {**{i: d["name"] for i, d in hand_items.items() if isinstance(d, dict)},
                                               **{i: d["name"] for i, d in items.items()}}.items() if v == b.group(1)), None)
                bonus = {"item": bonus_id, "chance": int(b.group(2)) / 100.0}
            tool = str(g["Tool"] or "None")
            nodes[node_id] = {
                "name": g["Node"], "skill": skill, "item": g["Item ID"], "yield": int(g["Yield"]),
                "crit_yield": int(g["Crit Yield"]), "min_skill": int(g["Min Skill"]), "max_level": int(g["Max Level"]),
                "base_success": pct(g["Base Success"]), "crit_chance": pct(g["Crit Chance"]),
                "respawn_seconds": float(g["Respawn (s)"]), "gather_seconds": float(g["Gather Time (s)"]),
                "tool": "" if tool == "None" else {"Miner's Pick": "prospecting", "Woodcutter's Axe": "woodworking"}[tool],
                "bonus": bonus, "status": g["Status"], "location": g["Location"],
            }
            if g["Item ID"] in items:
                items[g["Item ID"]]["description"] = g["Description"]
                items[g["Item ID"]]["value"] = 2 + int(g["Min Skill"]) // 3

    # Crafting tabs
    for skill_tab in CRAFT_SKILLS:
        skill = skill_tab.lower()
        for r in rows(wb[skill_tab]):
            rid, out = r["Recipe ID"], r["Output Item ID"]
            learned = str(r["Learned From"] or "")
            cost, currency = parse_cost(r["Scroll Cost"])
            recipe = {
                "name": r["Recipe"], "skill": skill, "output": out, "yield": int(r["Yield"]),
                "crit_yield": int(r["Crit Yield"]), "ingredients": parse_ingredients(r["Ingredient IDs"]),
                "min_skill": int(r["Min Skill"]), "max_level": int(r["Max Level"]),
                "base_success": pct(r["Base Success"]), "crit_chance": pct(r["Crit Chance"]),
                "stations": [s.strip() for s in str(r["Station IDs"]).split(",") if s.strip()],
                "innate": learned.lower().startswith("innate"), "craft_time": float(r["Craft Time (s)"]),
                "craft_message": r["Flavour Text"] or "", "status": r["Status"],
            }
            recipes[rid] = recipe

            # Output item (existing items.json entries are left alone)
            if out not in hand_items and out not in VENDOR_ITEMS:
                unsupported = []
                item = items.get(out) or base_item(out, r["Recipe"], r["Description"] or "")
                item["description"] = r["Description"] or item["description"]
                item["value"] = 4 + int(r["Min Skill"]) * 2
                apply_effect(item, out, r["Effect / Stats"], unsupported)
                if item["type"] == "misc" and item["slot"] == "none":
                    item["value"] = 3 + int(r["Min Skill"])
                items[out] = item
                if unsupported:
                    report.append(f"{out}: not yet supported by the engine -> {'; '.join(unsupported)}")

            # Recipe scroll for anything learnable from a scroll
            if "scroll" in learned.lower() and not recipe["innate"]:
                sid = "recipe_" + rid
                recipe["scroll"] = sid
                if sid in hand_items:
                    continue
                scroll = base_item(sid, f"Recipe: {r['Recipe']}",
                                   f"Teaches the {skill_tab} recipe for {r['Recipe']}. Requires {skill_tab} {r['Min Skill']}.",
                                   "scroll", cost or 10)
                scroll["icon"] = "res://Assets/icons/items/scroll.png"
                scroll["currency_type"] = currency
                scroll["stackable"] = False
                scroll["teaches_recipe"] = rid
                scroll["skill"] = skill
                items[sid] = scroll

    # Vendor-only items and kits
    for item_id, (name, item_type, value, desc, extra) in VENDOR_ITEMS.items():
        if item_id in hand_items:
            continue
        item = items.get(item_id) or base_item(item_id, name, desc, item_type, value)
        item.update({"name": name, "type": item_type, "value": value, "description": desc})
        item.update(extra)
        if item_type != "misc":
            item["stackable"] = False
        item["icon"] = icon_for(item_id, item.get("slot", "none"))
        items[item_id] = item
    for kit_id, skill in KIT_NAMES.items():
        if kit_id in hand_items:
            continue
        name = " ".join(w.capitalize() for w in kit_id.split("_"))
        kit = base_item(kit_id, name, f"A portable {skill} kit. Right-click and Open to craft anywhere. In a character-sheet slot it is also an "
                         f"8-slot bag for {skill} materials.", "tool", 25)
        kit.update({"stackable": False, "skill": skill, "weight": 3.0, "tradeskill_station": kit_id})
        if not os.path.exists(os.path.join(ICON_DIR, kit_id + ".png")):
            kit["icon"] = f"res://Assets/icons/items/{KIT_ICONS.get(kit_id, 'pouch.png')}"
        items[kit_id] = kit
    for meat_id, mob in MEATS.items():
        if meat_id in items:
            items[meat_id]["description"] = f"Raw {mob} meat. Cook it before you eat it."
            items[meat_id]["value"] = 1

    # Sanity: every ingredient and output resolves to an item
    known = set(hand_items) | set(items)
    for rid, r in recipes.items():
        for i in list(r["ingredients"]) + [r["output"]]:
            if i not in known and i != "cooked_meat_any":
                report.append(f"{rid}: unknown item '{i}'")
    for nid, n in nodes.items():
        if n["item"] not in known:
            report.append(f"node {nid}: unknown item '{n['item']}'")

    def write(name, comment, body):
        path = os.path.join(ROOT, "Data", name)
        with open(path, "w") as f:
            json.dump({"_comment": comment, **body}, f, indent=2, ensure_ascii=False)
            f.write("\n")

    # Ingredient groups: 'cooked_meat_any' in a recipe matches any cooked meat
    groups = {"cooked_meat_any": sorted(i for i in known if i.startswith("cooked_") and i.endswith("meat"))}

    src = "Generated by tools/export_crafting.py from Crafting.xlsx - do not edit by hand; edit the workbook and re-run."
    write("tradeskill_recipes.json",
          f"tradeskill_recipes.json - every crafting recipe keyed by recipe_id (Scripts/tradeskill_window.gd). "
          f"'stations' lists every station id the recipe can be made at; a group id in ingredients (e.g. cooked_meat_any) matches any "
          f"item listed under that key in 'ingredient_groups'. {src}",
          {"ingredient_groups": groups, "recipes": recipes})
    write("gathering_nodes.json",
          f"gathering_nodes.json - gatherable node types (Scripts/gathering_node.gd). 'tool' is the gather_tool an item "
          f"must have in your bags ('' = none). Where they are placed lives in Data/gathering_node_placements.json. {src}",
          {"nodes": nodes})
    # New items go on the end of items.json (same formatting as the file already has: 2-space indent, UTF-8)
    new_items = {k: v for k, v in items.items() if k not in hand_items}
    if new_items:
        hand_items.update(new_items)
        with open(items_path, "w") as f:
            f.write(json.dumps(hand_items, indent=2, ensure_ascii=False) + "\n")

    print(f"Exported {len(recipes)} recipes, {len(nodes)} gathering nodes; added {len(new_items)} new item(s) to "
          f"Data/items.json{': ' + ', '.join(new_items) if new_items else ''}.")
    if report:
        print("\nNotes:")
        for line in report:
            print("  - " + line)


if __name__ == "__main__":
    main()
